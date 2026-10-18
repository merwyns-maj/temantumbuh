#!/usr/bin/env python3
"""Build the Kitab TemanTumbuh workbook from kitab/kitab.json.

Usage: python kitab/build_kitab.py [out.xlsx]
Upload the .xlsx to the TemanTumbuh Google Drive and open it as a Google Sheet (File > Save as Google Sheets).
Colours and type follow design/tokens.json; Nunito is a Google font, so Sheets renders it.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
K = json.load(open(os.path.join(ROOT, "kitab", "kitab.json"), encoding="utf-8"))
TOK = {t["name"]: t["value"] for t in json.load(open(os.path.join(ROOT, "design", "tokens.json")))["color"]["tokens"]}


def c(name):
    v = TOK[name]
    v = v["light"] if isinstance(v, dict) else v
    if v.startswith("{"):
        return c(v.strip("{}"))
    return v.lstrip("#").upper()


FONT = "Nunito"
SRC_URL = {
    "AAP, A Parent's Guide to Safe Sleep (HealthyChildren.org)": "https://www.healthychildren.org/English/ages-stages/baby/sleep/Pages/a-parents-guide-to-safe-sleep.aspx",
    "AAP, Car Seats: Information for Families (HealthyChildren.org)": "https://www.healthychildren.org/English/safety-prevention/on-the-go/Pages/Car-Safety-Seats-Information-for-Families.aspx",
    "AAP, Fever: When to Call the Pediatrician (HealthyChildren.org)": "https://www.healthychildren.org/English/health-issues/conditions/fever/Pages/When-to-Call-the-Pediatrician.aspx",
    "IDAI, Penggunaan Baby Walker": "https://www.idai.or.id/artikel/seputar-kesehatan-anak/penggunaan-baby-walker",
    "Orlando Health, Is Babywearing Hurting My Child's Hips? (2019)": "https://www.orlandohealth.com/content-hub/is-babywearing-hurting-my-childs-hips",
    "IDAI, Keterlambatan Bicara": "https://www.idai.or.id/artikel/klinik/keluhan-anak/keterlambatan-bicara",
    "IDAI, Mencegah Terlambat Bicara pada Anak": "https://www.idai.or.id/artikel/klinik/pengasuhan-anak/mencegah-terlambat-bicara-pada-anak",
    "BPOM": "https://cekbpom.pom.go.id",
    "Kemenkes RI, via Antara: Jangan sepelekan buku KIA": "https://www.antaranews.com/berita/2206862/jangan-sepelekan-buku-kia-untuk-pantau-tumbuh-kembang-anak",
}
CDC = "https://www.cdc.gov/act-early/milestones/index.html"

thin = Side(style="thin", color=c("line"))
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
TIER = {"WAJIB": ("sage-ink", "on-accent"), "BERGUNA": ("seri-mpasi-bg", "seri-mpasi-ink"), "SKIP DULU": ("blush", "terra-ink")}


def fill(name):
    return PatternFill("solid", start_color=c(name), end_color=c(name))


def title(ws, text, sub=None):
    ws.sheet_view.showGridLines = False
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=20, bold=True, color=c("ink"))
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, size=11, color=c("ink-muted"))


def table(ws, top, headers, rows, widths, tier_col=None, link_cols=(), src_col=None, todo_cols=()):
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=top, column=j, value=h)
        cell.font = Font(name=FONT, size=11, bold=True, color=c("on-accent"))
        cell.fill = fill("deep-block")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[top].height = 32
    for i, r in enumerate(rows, top + 1):
        tallest = 1
        for j, v in enumerate(r, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name=FONT, size=10, color=c("ink"))
            cell.alignment = WRAP
            cell.border = BORDER
            cell.fill = fill("surface-card") if i % 2 else fill("surface-100")
            skip_row = tier_col is not None and str(r[tier_col]).startswith("SKIP")
            if j - 1 in todo_cols and not v and not skip_row:
                cell.value = "Segera"
                cell.font = Font(name=FONT, size=10, italic=True, color=c("dot-inactive"))
            if tier_col is not None and j - 1 == tier_col:
                key = next((t for t in TIER if str(v).startswith(t)), None)
                if key:
                    bg, fg = TIER[key]
                    cell.fill = fill(bg)
                    cell.font = Font(name=FONT, size=10, bold=True, color=c(fg))
            if src_col is not None and j - 1 == src_col and v:
                url = SRC_URL.get(v) or (CDC if v.startswith("CDC") else None)
                if url:
                    cell.hyperlink = url
                    cell.font = Font(name=FONT, size=10, color=c("sage-ink"), underline="single")
            if j - 1 in link_cols and v and str(v).startswith("http"):
                cell.hyperlink = v
                cell.font = Font(name=FONT, size=10, color=c("terra-deep"), underline="single", bold=True)
            w = widths[j - 1]
            tallest = max(tallest, -(-len(str(cell.value or "")) // max(1, int(w * 1.1))))
        ws.row_dimensions[i].height = max(30, 15 * tallest + 6)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = ws.cell(row=top + 1, column=3)
    return top + len(rows)


def main(out):
    wb = Workbook()
    # ---------------- Baca Dulu
    ws = wb.active
    ws.title = "Baca Dulu"
    cv = K["cover"]
    title(ws, cv["title"], cv["tagline"])
    ws["A3"] = cv["updated"]
    ws["A3"].font = Font(name=FONT, size=10, italic=True, color=c("ink-muted"))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws.sheet_view.showGridLines = False
    r = 5
    for head, lines in cv["sections"]:
        ws.cell(row=r, column=1, value=head).font = Font(name=FONT, size=12, bold=True, color=c("sage-ink"))
        if any(l.startswith(("WAJIB:", "BERGUNA:", "SKIP DULU:")) for l in lines):
            r += 1
        for line in lines:
            is_tier = line.startswith(("WAJIB:", "BERGUNA:", "SKIP DULU:"))
            cell = ws.cell(row=r, column=2, value=line.split(": ", 1)[1] if is_tier else line)
            cell.font = Font(name=FONT, size=11, color=c("ink"))
            cell.alignment = WRAP
            ws.row_dimensions[r].height = max(18, 16 * -(-len(line) // 105))
            if is_tier:
                key = line.split(":")[0]
                bg, fg = TIER[key]
                ws.cell(row=r, column=1).fill = fill(bg)
                ws.cell(row=r, column=1).value = key
                ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True, color=c(fg))
            r += 1
        r += 1
    # ---------------- Belanja Newborn
    ws = wb.create_sheet("Belanja Newborn")
    title(ws, "Belanja Newborn 0-3 bulan", "Wajib, berguna, atau skip dulu. Pilihan produk dan link toko segera diisi.")
    rows = [[k, b, t, why, cek, "", "", "", note, src] for k, b, t, why, cek, note, src in K["newborn"]]
    table(ws, 4, ["Kategori", "Barang", "Tier", "Kenapa", "Cek sebelum beli", "Pilihan kami", "Shopee", "TikTok Shop", "Catatan", "Sumber"],
          rows, [16, 26, 14, 36, 36, 20, 12, 12, 34, 30], tier_col=2, src_col=9, todo_cols=(5, 6, 7))
    # ---------------- Worksheet
    ws = wb.create_sheet("Worksheet 1-5 Tahun")
    title(ws, "Worksheet dan buku aktivitas 1-5 tahun", "Sesuai yang biasanya udah bisa di tiap usia. Tiap anak punya waktunya sendiri.")
    rows = [[u, bisa, akt, skip, cek, "", "", "", "", src] for u, bisa, akt, skip, cek, src in K["worksheet"]]
    table(ws, 4, ["Usia", "Biasanya udah bisa", "Aktivitas kertas yang pas", "Skip dulu", "Cek sebelum beli", "Pilihan buku aktivitas", "Shopee", "TikTok Shop", "Versi TemanTumbuh", "Sumber"],
          rows, [16, 34, 36, 26, 28, 22, 12, 12, 18, 32], src_col=9, todo_cols=(5, 6, 7, 8))
    # ---------------- Mainan Ngajak Ngobrol
    ws = wb.create_sheet("Mainan Ngajak Ngobrol")
    title(ws, "Speech delay? Mainan yang ngajak ngobrol", K["speech_note"])
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 34
    ws["A3"] = K["speech_screen"]
    ws["A3"].font = Font(name=FONT, size=11, bold=True, color=c("sage-ink"))
    ws.merge_cells("A3:I3")
    ws["A5"] = "Tanda yang perlu dibawa ke dokter anak (IDAI)"
    ws["A5"].font = Font(name=FONT, size=12, bold=True, color=c("terra-ink"))
    ws["A5"].hyperlink = SRC_URL["IDAI, Keterlambatan Bicara"]
    rr = 6
    for usia, tanda in K["speech_flags"]:
        a = ws.cell(row=rr, column=1, value=usia)
        b = ws.cell(row=rr, column=2, value=tanda)
        for cell in (a, b):
            cell.fill = fill("blush")
            cell.border = BORDER
            cell.alignment = WRAP
        a.font = Font(name=FONT, size=10, bold=True, color=c("terra-ink"))
        b.font = Font(name=FONT, size=10, color=c("ink"))
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
        ws.row_dimensions[rr].height = 30
        rr += 1
    top = rr + 1
    rows = [[u, m, why, cara, cek, "", "", "", src] for u, m, why, cara, cek, src in K["speech"]]
    table(ws, top, ["Usia", "Mainan atau aktivitas", "Kenapa bantu ngobrol", "Cara mainnya", "Cek sebelum beli", "Pilihan kami", "Shopee", "TikTok Shop", "Sumber"],
          rows, [16, 30, 32, 42, 30, 20, 12, 12, 30], src_col=8, todo_cols=(5, 6, 7))
    ws.freeze_panes = None
    # ---------------- Gratisan
    ws = wb.create_sheet("Gratisan")
    title(ws, "Gratisan dari TemanTumbuh", "Buat kamu, gratis.")
    table(ws, 4, ["Nama", "Isinya", "Link", "Cara pakai"], K["gratisan"], [30, 44, 30, 60], link_cols=(2,))
    ws.freeze_panes = None
    for s in wb.worksheets:
        s.sheet_properties.tabColor = c("terra") if s.title == "Baca Dulu" else c("sage")
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "out", "Kitab-TemanTumbuh.xlsx"))
